import os


def extract_text_from_file(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        import fitz
        doc = fitz.open(file_path)
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
        return text

    elif ext in (".pptx", ".ppt"):
        from pptx import Presentation
        prs = Presentation(file_path)
        lines = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
        return "\n".join(lines)

    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()


ALLOWED_EXTENSIONS = {".pdf", ".pptx", ".ppt", ".txt"}


def validate_file_extension(filename: str) -> bool:
    ext = os.path.splitext(filename)[1].lower()
    return ext in ALLOWED_EXTENSIONS
